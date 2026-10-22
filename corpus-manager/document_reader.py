import os
import re
from typing import List, Optional, Dict, Any
from pathlib import Path

import sys
sys.path.append(os.path.join(os.path.dirname(__file__), '..', 'LAB1'))
from text_io import extract_text_from_txt, extract_text_from_rtf

try:
    import PyPDF2
    PYPDF2_AVAILABLE = True
except ImportError:
    PYPDF2_AVAILABLE = False

try:
    import docx
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False

try:
    import win32com.client
    WIN32COM_AVAILABLE = True
except ImportError:
    WIN32COM_AVAILABLE = False


class DocumentReader:
    
    def __init__(self):
        self.supported_formats = {
            '.txt': self._read_txt,
            '.rtf': self._read_rtf,
        }

        if PYPDF2_AVAILABLE:
            self.supported_formats['.pdf'] = self._read_pdf

        if DOCX_AVAILABLE:
            self.supported_formats['.docx'] = self._read_docx

        if WIN32COM_AVAILABLE:
            self.supported_formats['.doc'] = self._read_doc
    
    def get_supported_formats(self) -> List[str]:
        return list(self.supported_formats.keys())
    
    def is_format_supported(self, file_path: str) -> bool:
        ext = Path(file_path).suffix.lower()
        return ext in self.supported_formats
    
    def read_document(self, file_path: str) -> str:
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"Файл не найден: {file_path}")
        
        ext = Path(file_path).suffix.lower()
        
        if ext not in self.supported_formats:
            raise ValueError(f"Неподдерживаемый формат файла: {ext}")
        
        try:
            return self.supported_formats[ext](file_path)
        except Exception as e:
            raise RuntimeError(f"Ошибка при чтении файла {file_path}: {e}")
    
    def _read_txt(self, file_path: str) -> str:
        return extract_text_from_txt(file_path)
    
    def _read_rtf(self, file_path: str) -> str:
        return extract_text_from_rtf(file_path)
    
    def _read_pdf(self, file_path: str) -> str:
        if not PYPDF2_AVAILABLE:
            raise RuntimeError("PyPDF2 не доступен. Установите: pip install PyPDF2")
        
        text = []

        try:
            with open(file_path, 'rb') as file:
                pdf_reader = PyPDF2.PdfReader(file)
                
                for page_num in range(len(pdf_reader.pages)):
                    page = pdf_reader.pages[page_num]
                    page_text = page.extract_text()

                    page_text = self._clean_pdf_text(page_text)
                    text.append(page_text)
        
        except Exception as e:
            raise RuntimeError(f"Ошибка при чтении PDF: {e}")
        
        return "\n\n".join(text)
    
    def _clean_pdf_text(self, text: str) -> str:
        text = re.sub(r'\s+', ' ', text)

        text = re.sub(r'(\w+)-\s+(\w+)', r'\1\2', text)

        text = re.sub(r'[^\w\s\u0400-\u04FF.,!?;:()\-"\'\n]', '', text)
        
        return text.strip()
    
    def _read_docx(self, file_path: str) -> str:
        if not DOCX_AVAILABLE:
            raise RuntimeError("python-docx не доступен. Установите: pip install python-docx")
        
        try:
            doc = docx.Document(file_path)
            paragraphs = []

            for paragraph in doc.paragraphs:
                if paragraph.text.strip():
                    paragraphs.append(paragraph.text)

            for table in doc.tables:
                for row in table.rows:
                    row_text = []
                    for cell in row.cells:
                        if cell.text.strip():
                            row_text.append(cell.text.strip())
                    if row_text:
                        paragraphs.append(" | ".join(row_text))
            
            return "\n\n".join(paragraphs)

        except Exception as e:
            raise RuntimeError(f"Ошибка при чтении DOCX: {e}")
    
    def _read_doc(self, file_path: str) -> str:
        if not WIN32COM_AVAILABLE:
            raise RuntimeError("pywin32 не доступен. Установите: pip install pywin32")
        
        if os.name != 'nt':
            raise RuntimeError("Чтение DOC файлов поддерживается только на Windows")

        try:
            temp_txt = file_path + ".temp.txt"

            word = win32com.client.Dispatch("Word.Application")
            word.Visible = False
            
            doc = word.Documents.Open(os.path.abspath(file_path))

            doc.SaveAs(os.path.abspath(temp_txt), FileFormat=2)
            
            doc.Close()
            word.Quit()

            try:
                with open(temp_txt, 'r', encoding='utf-8') as f:
                    text = f.read()
            except UnicodeDecodeError:
                with open(temp_txt, 'r', encoding='cp1251') as f:
                    text = f.read()

            try:
                os.remove(temp_txt)
            except:
                pass
            
            return text
        
        except Exception as e:
            raise RuntimeError(f"Ошибка при чтении DOC: {e}")
    
    def batch_read_documents(self, directory: str,
                           file_patterns: List[str] = None,
                           recursive: bool = False) -> Dict[str, str]:
        if file_patterns is None:
            file_patterns = ['*.txt', '*.rtf']
        
        directory_path = Path(directory)
        if not directory_path.exists():
            raise FileNotFoundError(f"Директория не найдена: {directory}")
        
        documents = {}

        if recursive:
            files_found = []
            for pattern in file_patterns:
                files_found.extend(directory_path.rglob(pattern))
        else:
            files_found = []
            for pattern in file_patterns:
                files_found.extend(directory_path.glob(pattern))
        
        for file_path in files_found:
            if self.is_format_supported(str(file_path)):
                try:
                    text = self.read_document(str(file_path))
                    documents[str(file_path)] = text
                except Exception as e:
                    print(f"Ошибка при чтении файла {file_path}: {e}")
        
        return documents
    
    def get_document_metadata(self, file_path: str) -> Dict[str, Any]:
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"Файл не найден: {file_path}")
        
        file_path_obj = Path(file_path)

        metadata = {
            'filename': file_path_obj.name,
            'filepath': str(file_path_obj.absolute()),
            'extension': file_path_obj.suffix.lower(),
            'size': file_path_obj.stat().st_size,
            'modified_time': file_path_obj.stat().st_mtime,
            'is_supported': self.is_format_supported(file_path)
        }

        ext = file_path_obj.suffix.lower()
        
        if ext == '.pdf' and PYPDF2_AVAILABLE:
            try:
                with open(file_path, 'rb') as file:
                    pdf_reader = PyPDF2.PdfReader(file)
                    metadata['pdf_pages'] = len(pdf_reader.pages)
                    
                    if pdf_reader.metadata:
                        metadata['pdf_title'] = pdf_reader.metadata.get('/Title', '')
                        metadata['pdf_author'] = pdf_reader.metadata.get('/Author', '')
                        metadata['pdf_creator'] = pdf_reader.metadata.get('/Creator', '')
            except:
                pass
        
        elif ext == '.docx' and DOCX_AVAILABLE:
            try:
                doc = docx.Document(file_path)
                metadata['docx_paragraphs'] = len(doc.paragraphs)
                metadata['docx_tables'] = len(doc.tables)
                
                core_props = doc.core_properties
                if core_props:
                    metadata['docx_title'] = core_props.title or ''
                    metadata['docx_author'] = core_props.author or ''
                    metadata['docx_created'] = core_props.created.isoformat() if core_props.created else ''
            except:
                pass
        
        return metadata
    
    def estimate_reading_time(self, text: str, words_per_minute: int = 200) -> Dict[str, Any]:
        words = len(text.split())
        characters = len(text)
        
        reading_time_minutes = words / words_per_minute
        reading_time_hours = reading_time_minutes / 60
        
        return {
            'words': words,
            'characters': characters,
            'reading_time_minutes': round(reading_time_minutes, 2),
            'reading_time_hours': round(reading_time_hours, 2),
            'estimated_complexity': self._estimate_complexity(text)
        }
    
    def _estimate_complexity(self, text: str) -> str:
        words = text.split()
        if not words:
            return "пустой"

        avg_word_length = sum(len(word.strip('.,!?;:()"\'')) for word in words) / len(words)

        sentences = re.split(r'[.!?]+', text)
        sentences = [s.strip() for s in sentences if s.strip()]
        
        if sentences:
            avg_sentence_length = sum(len(s.split()) for s in sentences) / len(sentences)
        else:
            avg_sentence_length = 0
        if avg_word_length < 4 and avg_sentence_length < 10:
            return "очень легкий"
        elif avg_word_length < 5 and avg_sentence_length < 15:
            return "легкий"
        elif avg_word_length < 6 and avg_sentence_length < 20:
            return "средний"
        elif avg_word_length < 7 and avg_sentence_length < 25:
            return "сложный"
        else:
            return "очень сложный"


_reader_instance = None


def get_document_reader() -> DocumentReader:
    global _reader_instance
    if _reader_instance is None:
        _reader_instance = DocumentReader()
    return _reader_instance


def load_text_from_file(file_path: str) -> str:
    reader = get_document_reader()
    return reader.read_document(file_path)
